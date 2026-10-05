"""embedded-list: pure helpers, one company end to end with fake pages, resume, workbook. No network."""

import asyncio

import pytest

from agent import embedded_list as el
from agent import web

HOME = """<html><body><h1>Acme Embedded</h1>
<p>We do Yocto and embedded Linux, BSP development and board bring-up, OTA updates and secure boot.</p>
<a href="/contact">Contact us</a> <a href="/about">About</a> <a href="/careers">Careers</a>
<a href="https://www.linkedin.com/company/acme">LinkedIn</a></body></html>"""
CONTACT = """<html><body>Acme Embedded GmbH, Hauptstr. 1, Munich, Germany. Offices in India and France.
Write to <a href="mailto:info@acme-embedded.de">info@acme-embedded.de</a> or careers@acme-embedded.de.
Our CEO: jane.doe@acme-embedded.de. <form><textarea></textarea></form></body></html>"""
ABOUT = "<html><body>Founded in Germany. Linux kernel and device driver experts.</body></html>"
CAREERS = """<html><body><h1>Join us</h1>
<a href="/careers/embedded-linux-engineer">Senior Embedded Linux Engineer (BSP)</a>
<a href="/careers/android-app-developer">Android App Developer</a>
<a href="/careers/closed-firmware">Firmware Engineer</a>
<a href="/careers">All jobs</a></body></html>"""
JOB_OK = """<html><body>Senior Embedded Linux Engineer. Location: Munich, Germany   Type: Full-time.
Hybrid work. You have 5+ years of experience with Yocto and U-Boot.</body></html>"""
JOB_CLOSED = "<html><body>Firmware Engineer. This position has been filled.</body></html>"

PAGES = {
    "https://acme-embedded.de": HOME,
    "https://acme-embedded.de/contact": CONTACT,
    "https://acme-embedded.de/about": ABOUT,
    "https://acme-embedded.de/careers": CAREERS,
    "https://acme-embedded.de/careers/embedded-linux-engineer": JOB_OK,
    "https://acme-embedded.de/careers/closed-firmware": JOB_CLOSED,
}


@pytest.fixture
def fake_web(monkeypatch):
    fetched: list[str] = []

    async def fetch_smart(url):
        fetched.append(url)
        return PAGES.get(url.rstrip("/"))

    monkeypatch.setattr(web, "fetch_smart", fetch_smart)
    return fetched


def test_links_are_absolute_and_unique():
    got = el.links(
        '<a href="/a">A</a><a href="/a#x">A again</a><a href="mailto:x@y.z">m</a>', "https://s.io/"
    )
    assert got == [("https://s.io/a", "A")]


def test_focus_tags():
    assert el.focus_tags("Yocto, BSP and OTA for Android Automotive (AAOS)") == [
        "aaos",
        "yocto",
        "ota",
        "bsp",
    ]
    assert el.focus_tags("We build web shops") == []


def test_company_emails_are_role_addresses_on_own_domain_careers_first():
    text = "info@acme.io careers@acme.io jane.doe@acme.io sales@other.com noreply@acme.io"
    assert el.company_emails(text, {"acme.io"}) == [
        "careers@acme.io",
        "info@acme.io",
    ]  # no person, no other domain


def test_relevance_of_job_titles():
    for title in ("Embedded Linux Engineer", "Firmware Developer", "Android Platform Engineer", "BSP Engineer",
                  "Yocto Build Engineer", "Linux Kernel Developer", "Embedded Security Engineer"):  # fmt: skip
        assert el.is_relevant(title), title
    for title in (
        "Android App Developer",
        "Sales Manager",
        "Frontend Developer",
        "Linux System Administrator",
    ):
        assert not el.is_relevant(title), title


def test_job_details_and_closed():
    d = el.job_details(
        "Location: Pune, India   Department: R&D. Remote possible. 3-5 years of experience."
    )
    assert d == {"location": "Pune, India", "remote": "Remote", "experience": "3-5 years"}
    assert el.is_closed("Sorry, this position has been filled.")
    assert not el.is_closed("Apply now")


def test_guess_country_prefers_site_ending_when_named():
    assert (
        el.guess_country(["Munich, Germany. Offices in India and France. Germany."], "acme.de")[0]
        == "Germany"
    )
    hq, others = el.guess_country(["HQ in Pune, India. Sales office in the USA."], "acme.com")
    assert hq == "India" and "United States" in others
    assert el.guess_country([""], "acme.com.br")[0] == "Brazil"


def test_dedupe_by_domain_alias_and_name():
    cands = [
        el.Candidate("Northern.tech (Mender)", "northern.tech", aliases=["mender.io"]),
        el.Candidate("Mender", "mender.io"),
        el.Candidate("Acme GmbH", "acme.de"),
        el.Candidate("ACME", "acme-other.com"),
    ]
    assert [c.name for c in el.dedupe(cands)] == ["Northern.tech (Mender)", "Acme GmbH"]


def test_research_one_company_end_to_end(fake_web):
    row, jobs = asyncio.run(
        el.research(el.Candidate("Acme Embedded", "acme-embedded.de"), "2026-10-05")
    )
    assert row.accessible and row.fits
    assert (
        row.contact_email == "careers@acme-embedded.de"
    )  # careers before info, never the CEO's address
    assert row.email_source_url == "https://acme-embedded.de/contact"
    assert row.careers_page == "https://acme-embedded.de/careers"
    assert row.country == "Germany" and "India" in row.other_countries
    assert {"yocto", "bsp", "ota", "security", "kernel"} <= set(row.focus_areas.split(", "))
    assert [j.job_title for j in jobs] == [
        "Senior Embedded Linux Engineer (BSP)"
    ]  # app job not relevant, closed one skipped
    assert (
        jobs[0].experience == "5+ years"
        and jobs[0].remote == "Hybrid"
        and "Munich" in jobs[0].location
    )
    assert not any("linkedin" in u for u in fake_web)  # portals are never opened


def test_unreachable_site_is_reported_not_guessed(fake_web):
    row, jobs = asyncio.run(el.research(el.Candidate("Gone Ltd", "gone.example"), "2026-10-05"))
    assert not row.accessible and row.notes.startswith(el.CANNOT_ACCESS) and not jobs
    assert row.contact_email == ""


def test_run_resumes_and_builds_the_workbook(fake_web, monkeypatch, tmp_path):
    from openpyxl import load_workbook

    cfg = {"seeds": [{"name": "Acme Embedded", "website": "acme-embedded.de"},
                     {"name": "Gone Ltd", "website": "gone.example"}]}  # fmt: skip
    monkeypatch.setattr(el, "load_config", lambda path=el.CONFIG: cfg)
    report = asyncio.run(el.run(discover=False, out=tmp_path))
    assert report["companies"] == 2 and report["total_openings"] == 1
    assert report["could_not_access"] == ["Gone Ltd"]
    asyncio.run(el.run(discover=False, out=tmp_path))  # second run: everything already done
    progress = (tmp_path / "companies_progress.csv").read_text().splitlines()
    assert len(progress) == 3  # header + 2 companies: nothing researched twice

    wb = load_workbook(tmp_path / "embedded_companies_and_openings.xlsx")
    assert wb.sheetnames == ["Companies", "Openings", "Summary"]
    ws = wb["Companies"]
    assert ws.freeze_panes == "A2" and ws.auto_filter.ref and ws["A1"].font.bold
    assert ws["B2"].hyperlink is not None
    assert (
        wb["Openings"].max_row == 2
    )  # the one opening (Gone Ltd could not be accessed: no placeholder row)
    assert (tmp_path / "companies_progress.csv").exists() and (
        tmp_path / "verification.csv"
    ).exists()


def test_verify_flags_email_without_source():
    rows = [
        el.CompanyRow(
            name="X", website="https://x.io", contact_email="info@x.io", email_source_url=""
        )
    ]
    checks = asyncio.run(el.verify(rows, [], seed=1))
    assert any(c["why"] == "email has no source URL" for c in checks)


def test_obfuscated_published_addresses_are_read():
    assert el.company_emails("write to info [at] acme [dot] io", {"acme.io"}) == ["info@acme.io"]
    assert el.company_emails("jobs(at)acme.io", {"acme.io"}) == ["jobs@acme.io"]


def test_country_from_address_beats_mere_mentions():
    text = "We spoke in Spain, Spain and Spain at conferences. Registered office: 12 rue X, 31000 Toulouse, France."
    assert el.guess_country([text], "bootlin.com")[0] == "France"


def test_plural_positions_link_counts_as_opening():
    page = [
        (
            "https://bootlin.com/company/kernel-embedded-linux-engineers",
            "Embedded Linux and Linux kernel engineers positions",
        )
    ]
    got = el.job_links(page, "https://bootlin.com/company/careers", {"bootlin.com"})
    assert got and el.is_relevant(got[0][1])


def test_relative_links_resolve_against_the_canonical_address():
    html = (
        '<link rel="canonical" href="https://bootlin.com/company/careers/" />'
        '<a href="kernel-engineers">Kernel engineers positions</a>'
    )
    assert el.links(html, "https://bootlin.com/company/careers") == [
        ("https://bootlin.com/company/careers/kernel-engineers", "Kernel engineers positions")
    ]


def test_staff_and_blog_pages_are_never_info_pages():
    page = [
        ("https://bootlin.com/company/staff/paul/", "Paul"),
        ("https://bootlin.com/blog/x/", "About our talk"),
        ("https://bootlin.com/company/", "Company"),
        ("https://bootlin.com/contact-us/", "Get in touch"),
    ]
    assert el.info_links(page, {"bootlin.com"}) == [
        "https://bootlin.com/contact-us/",
        "https://bootlin.com/company/",
    ]


def test_view_openings_link_is_followed(fake_web, monkeypatch):
    pages = {
        "https://mistral.example": '<a href="/career/">Careers</a> embedded Linux BSP',
        "https://mistral.example/career": "<h4>Senior V&V Engineer</h4>"  # a testimonial, not an opening
        '<a href="/careers-job-listings/">View Openings</a>',
        "https://mistral.example/careers-job-listings": '<a href="/jobs/42">Embedded Firmware Engineer</a>',
        "https://mistral.example/jobs/42": "Embedded Firmware Engineer. Location: Bengaluru, India   Type: full time",
    }
    for k, v in pages.items():
        monkeypatch.setitem(PAGES, k, v)
    _, jobs = asyncio.run(el.research(el.Candidate("Mistral", "mistral.example"), "2026-10-05"))
    assert [j.job_title for j in jobs] == ["Embedded Firmware Engineer"]
    assert jobs[0].location == "Bengaluru, India"


def test_canonical_found_after_a_huge_head():
    html = (
        "<head><style>"
        + "x" * 50_000
        + '</style><link rel="canonical" href="https://s.io/careers/" /></head>'
    )
    assert el.page_base(html, "https://s.io/careers") == "https://s.io/careers/"


def test_job_listing_navigation_link_is_not_a_job():
    page = [
        ("https://m.io/career/careers-job-listings/", "Careers Job Listings"),
        ("https://m.io/jobs/1", "View Openings"),
    ]
    assert el.job_links(page, "https://m.io/career/", {"m.io"}) == []


def test_view_details_links_take_the_title_from_the_document_name():
    base = "https://m.io/wp-content/uploads/2026/07/"
    page = [
        (
            base + "13.Embedded-SoftwareLinux-Vxworks-Senior-Engineer-_Module-Lead.docx",
            "View Details",
        ),
        (base + "Production-Manager-Project-Lead.docx", "View Details"),
        ("https://docs.google.com/forms/d/e/x/viewform", "Apply"),  # another site: never an opening
        ("https://m.io/about-us/culture/", "Learn more"),
    ]
    got = el.job_links(page, "https://m.io/careers-job-listings/", {"m.io"})
    assert [t for _, t in got] == [
        "Embedded SoftwareLinux Vxworks Senior Engineer Module Lead",
        "Production Manager Project Lead",
    ]
    assert el.is_relevant(got[0][1]) and not el.is_relevant(got[1][1])


def test_found_site_gets_its_name_from_its_own_page():
    og = (
        '<meta property="og:site_name" content="Acme Embedded" /><title>BSP services | Acme</title>'
    )
    assert el.site_name(og, "https://acme-embedded.com") == "Acme Embedded"
    assert (
        el.site_name(
            "<title>Android BSP Development | Acme Embedded</title>", "https://acme-embedded.com"
        )
        == "Acme Embedded"
    )
    assert el.site_name("<title>Home</title>", "https://acme-embedded.com") == "Acme-embedded"


def test_found_sites_need_to_be_opened_fit_and_not_skipped(monkeypatch, tmp_path):
    rows = [
        el.CompanyRow(
            name="Seed Down", website="https://down.io", accessible=False, found_via="seed"
        ),
        el.CompanyRow(
            name="Found Down",
            website="https://jobs.example",
            accessible=False,
            found_via="search: x",
        ),
        el.CompanyRow(
            name="Found Training",
            website="https://teach.example",
            focus_areas="bsp",
            found_via="search: x",
        ),
        el.CompanyRow(
            name="Found Good",
            website="https://good.example",
            focus_areas="bsp, yocto",
            found_via="search: x",
        ),
    ]
    el._write_csv(tmp_path / "companies_progress.csv", rows, el.CompanyRow)
    monkeypatch.setattr(
        el, "load_config", lambda path=el.CONFIG: {"skip_domains": ["teach.example"]}
    )
    report = asyncio.run(el.finish(tmp_path, seed=1))
    assert report["companies"] == 2
    assert report["left_out_not_embedded"] == ["Found Down", "Found Training"]


EXCLUDED = ["indiamart", "alibaba", "amazon", "upwork", "justdial"]


def test_marketplaces_are_excluded_sites():
    for url in ("https://www.indiamart.com/acme", "dir.indiamart.com", "https://m.alibaba.com/x",
                "https://www.amazon.in/dp/1", "https://aws.amazon.com", "https://www.upwork.com/jobs/1"):  # fmt: skip
        assert el.excluded_site(url, EXCLUDED), url
    assert not el.excluded_site("https://www.acme-embedded.com", EXCLUDED)
    assert not el.excluded_site(
        "https://www.amazonite-devices.com", EXCLUDED
    )  # a word part, not a host part


def test_marketplace_search_results_are_ignored(monkeypatch):
    async def search(q, limit=8):
        return [("https://www.indiamart.com/rfid-reader/", "RFID reader - IndiaMART"),
                ("https://www.acme-bsp.com/services", "BSP services | Acme BSP")]  # fmt: skip

    async def no_page(url):
        return None

    monkeypatch.setattr(el.sources, "web_search", search)
    monkeypatch.setattr(el, "_get", no_page)
    cfg = {"seeds": [], "discovery_queries": ["bsp"], "directory_pages": ["https://dir.example/list"],
           "excluded_sites": EXCLUDED}  # fmt: skip
    got = asyncio.run(el.candidates(cfg))
    assert [c.website for c in got] == ["acme-bsp.com"]
    assert el.directory_status == [
        {"page": "https://dir.example/list", "readable": False, "candidates": 0}
    ]


def test_online_store_is_recognised():
    assert el.looks_like_store("Add to cart. Free shipping on orders. Checkout securely. Wishlist")
    assert not el.looks_like_store("We do Yocto BSP work. Contact us. Buy now our dev kit.")


def test_source_check_flags_rows_citing_a_marketplace():
    companies = [
        el.CompanyRow(
            name="Good", website="https://good.io", email_source_url="https://good.io/contact"
        ),
        el.CompanyRow(
            name="Bad", website="https://good2.io", careers_page="https://www.upwork.com/x"
        ),
    ]
    jobs = [
        el.OpeningRow(
            company="Good", job_title="BSP Engineer", job_link="https://www.indiamart.com/j"
        )
    ]
    bad = el.source_checks(companies, jobs, EXCLUDED)
    assert [(b["company"], b["item"]) for b in bad] == [
        ("Bad", "https://www.upwork.com/x"),
        ("Good", "https://www.indiamart.com/j"),
    ]
    checks = asyncio.run(el.verify([], [], seed=1, excluded=EXCLUDED))
    assert (
        checks[-1]["kind"] == "source" and checks[-1]["ok"]
    )  # the summary check is always reported


def test_contact_form_note_and_no_opening_text_follow_the_spec(fake_web, monkeypatch, tmp_path):
    from openpyxl import load_workbook

    monkeypatch.setitem(
        PAGES,
        "https://formco.example",
        "<p>Yocto BSP services</p><form><textarea></textarea></form>",
    )
    row, _ = asyncio.run(el.research(el.Candidate("FormCo", "formco.example"), "2026-10-05"))
    assert row.contact_email == f"Not listed {chr(0x2013)} contact form: https://formco.example"
    assert "only a contact form, no email listed" in row.notes
    el.build_workbook([row], [], [], tmp_path / "w.xlsx")
    ws = load_workbook(tmp_path / "w.xlsx")["Openings"]
    assert ws["B2"].value == "No relevant opening found (checked 2026-10-05)"


def test_tidy_name_and_verification_reads_obfuscated_addresses(monkeypatch):
    assert el.tidy_name("DH electronics: DH electronics") == "DH electronics"
    assert el.tidy_name("Acme: Embedded Linux") == "Acme: Embedded Linux"

    async def page(url):
        return ("", "Impressum. E-Mail: info[at]dh-electronics.com")

    monkeypatch.setattr(el, "_get", page)
    row = el.CompanyRow(name="DH", website="https://dh-electronics.com", contact_email="info@dh-electronics.com",
                        email_source_url="https://www.dh-electronics.com/impressum")  # fmt: skip
    checks = asyncio.run(el.verify([row], [], seed=1))
    assert [c["ok"] for c in checks if c["kind"] == "email"] == [True]
