"""Research list: embedded engineering companies worldwide, their public contact emails and current relevant job
openings (docs/embedded-companies-research.md). `python -m agent embedded-list`.

Kept apart from the lead pipeline: nothing here touches the database, drafts an email or reads the lead search's
exclusion lists. Everything goes to out/embedded_companies/: progress CSVs (resumable), a verification CSV and the
Excel workbook. Emails, careers pages and openings are taken from pages actually fetched, by code; nothing is
guessed, and portals whose terms forbid scraping (LinkedIn, Indeed, Naukri...) are never fetched (agent.web).
"""

import asyncio
import csv
import random
import re
from collections import Counter
from dataclasses import asdict, dataclass, field, fields
from datetime import date
from html import unescape
from html.parser import HTMLParser
from pathlib import Path
from urllib.parse import unquote, urljoin, urlparse

import yaml

from agentkit.config import env
from agentkit.log import get_logger

from . import contacts, sources, web
from .gui import leadscore

log = get_logger("agent.embedded_list")
CONFIG = Path("config/embedded_companies.yaml")
OUT = Path("out/embedded_companies")
NO_OPENING = "No relevant opening found"
CANNOT_ACCESS = "Could not access - check manually"

# --- what the pages say ------------------------------------------------------------------------------

FOCUS = {  # tag -> wording on a company's own pages that shows real work in that area
    "firmware": r"\bfirmware\b|embedded software",
    "aosp": r"\baosp\b|android open source project|android (platform|framework|bsp|porting|customi[sz]ation)",
    "aaos": r"android automotive|\baaos\b",
    "yocto": r"\byocto\b|openembedded|\bbuildroot\b",
    "embedded-linux": r"embedded linux",
    "ota": r"\bota\b|over[- ]the[- ]air",
    "security": r"secure boot|trustzone|\bop-?tee\b|cyber resilience act|\bcra\b compliance|device security|"
    r"embedded security|security hardening",
    "bsp": r"\bbsp\b|board support package|board bring[- ]?up",
    "hal": r"\bhal\b|hardware abstraction layer",
    "kernel": r"linux kernel|device drivers?|kernel (development|drivers?|porting)",
}
_FOCUS = {k: re.compile(v, re.I) for k, v in FOCUS.items()}

# Company-level mailboxes: careers first (the task prefers them), then general contact, then the rest.
_ROLE_RANK = (
    ("careers", "career", "jobs", "job", "hr", "recruit", "talent", "hiring", "people", "join"),
    ("info", "contact", "hello", "office", "enquir", "inquir", "mail", "team", "general"),
    ("sales", "business", "partner", "support"),
)

_CAREERS_LINK = re.compile(
    r"career|\bjobs?\b|join[- ]?us|work[- ]with[- ]us|we.re hiring|vacanc|karriere|stellen|emplois|"
    r"carri[eè]re|lavora|trabaja|open positions|openings",
    re.I,
)
_INFO_LINK = re.compile(
    r"contact|kontakt|about|impressum|imprint|company|legal|offices|locations", re.I
)
# Pages about one person or one article: never read for the company's address or country.
_NOT_INFO = re.compile(
    r"/(staff|team|people|person|author|blog|news|events?|posts?|articles?|tag|category)/", re.I
)
_OPENINGS_LINK = re.compile(
    r"(view|see|browse|all|current|open)\s+(current\s+|all\s+|open\s+)?(openings|positions|jobs|vacancies|roles)|"
    r"job (listings|openings|board)|open positions|current openings|vacancies",
    re.I,
)
# Applicant systems that host companies' job boards (allowed: they publish the company's own openings).
ATS_HOSTS = (
    "greenhouse.io", "lever.co", "workable.com", "personio.de", "personio.com", "recruitee.com",
    "smartrecruiters.com", "teamtailor.com", "bamboohr.com", "ashbyhq.com", "join.com", "breezy.hr",
    "zohorecruit.com", "zohorecruit.in", "freshteam.com", "darwinbox.in", "keka.com", "myworkdayjobs.com",
    "icims.com", "jobvite.com", "rippling-ats.com", "homerun.co", "softgarden.io", "jobs.ashbyhq.com",
)  # fmt: skip
_JOB_PATH = re.compile(
    r"/(jobs?|careers?|positions?|openings?|vacanc(y|ies)|stellen(angebote)?|offres?|job-offers?|o|j)/[^/?#]{3,}",
    re.I,
)
_ROLE_WORD = re.compile(
    r"\b(engineer|developer|architect|lead|specialist|intern(ship)?|consultant|programmer|expert|scientist|"
    r"technician|tester|integrator|werkstudent|praktikant|entwickler|ingenieur|ing[ée]nieur)s?\b",
    re.I,
)
_NAV_TEXT = re.compile(
    r"^(all |view |see |browse |open )?(jobs|careers|positions|openings|vacancies)$|apply now|^apply$|"
    r"read more|learn more",
    re.I,
)
RELEVANT = re.compile(
    r"embedded|firmware|kernel|\bbsp\b|board support|bring[- ]?up|\baosp\b|android (platform|framework|system|"
    r"os|bsp|automotive|hal|middleware)|yocto|openembedded|buildroot|\bota\b|over[- ]the[- ]air|secure boot|"
    r"trustzone|op-?tee|(embedded|device|product|iot) security|\bhal\b|device drivers?|driver develop|"
    r"linux (kernel|bsp|drivers?|platform|system software|firmware)|u-boot|bootloader|\brtos\b|zephyr|bare[- ]metal",
    re.I,
)
_CLOSED = re.compile(
    r"no longer (accepting|available|open)|position (has been|is) (filled|closed)|job (has )?expired|"
    r"this (job|position|vacancy) (is|has been) closed|applications? (are |is )?(now )?closed|"
    r"vacancy (has been )?filled",
    re.I,
)
_LOCATION = re.compile(
    r"(?:job location|work location|location|standort|lieu|based in)\s*[:\--]\s*([^\n|•]{2,80}?)(?=\s{2,}|$|"
    r"\b(department|type|experience|employment|team|salary|apply)\b)",
    re.I,
)
_EXPERIENCE = re.compile(
    r"(\d{1,2})\s*(\+|plus)?\s*(?:(?:-|-|to)\s*(\d{1,2}))?\s*(?:years|yrs)\b(?:[^.]{0,40}?experience)?",
    re.I,
)

COUNTRIES = {
    "India": ("india",), "United States": ("united states", "usa", "u.s.a."), "Germany": ("germany", "deutschland"),
    "France": ("france",), "United Kingdom": ("united kingdom", "england", "scotland"), "Italy": ("italy", "italia"),
    "Spain": ("spain", "españa"), "Netherlands": ("netherlands", "the netherlands"), "Belgium": ("belgium",),
    "Switzerland": ("switzerland", "schweiz", "suisse"), "Austria": ("austria", "österreich"),
    "Poland": ("poland", "polska"), "Czech Republic": ("czech republic", "czechia"), "Romania": ("romania",),
    "Hungary": ("hungary",), "Portugal": ("portugal",), "Sweden": ("sweden", "sverige"), "Norway": ("norway",),
    "Denmark": ("denmark",), "Finland": ("finland",), "Ireland": ("ireland",), "Greece": ("greece",),
    "Ukraine": ("ukraine",), "Serbia": ("serbia",), "Croatia": ("croatia",), "Slovenia": ("slovenia",),
    "Estonia": ("estonia",), "Lithuania": ("lithuania",), "Latvia": ("latvia",), "Bulgaria": ("bulgaria",),
    "Turkey": ("turkey", "türkiye"), "Israel": ("israel",), "United Arab Emirates": ("united arab emirates", "uae"),
    "Canada": ("canada",), "Mexico": ("mexico",), "Brazil": ("brazil", "brasil"), "Argentina": ("argentina",),
    "Chile": ("chile",), "Colombia": ("colombia",), "China": ("china",), "Taiwan": ("taiwan",),
    "Japan": ("japan",), "South Korea": ("south korea", "korea"), "Singapore": ("singapore",),
    "Malaysia": ("malaysia",), "Vietnam": ("vietnam",), "Philippines": ("philippines",), "Indonesia": ("indonesia",),
    "Thailand": ("thailand",), "Australia": ("australia",), "New Zealand": ("new zealand",),
    "South Africa": ("south africa",), "Egypt": ("egypt",), "Pakistan": ("pakistan",), "Sri Lanka": ("sri lanka",),
}  # fmt: skip
_TLD_COUNTRY = {
    "in": "India", "de": "Germany", "fr": "France", "uk": "United Kingdom", "it": "Italy", "es": "Spain",
    "nl": "Netherlands", "be": "Belgium", "ch": "Switzerland", "at": "Austria", "pl": "Poland", "cz": "Czech Republic",
    "ro": "Romania", "hu": "Hungary", "pt": "Portugal", "se": "Sweden", "no": "Norway", "dk": "Denmark",
    "fi": "Finland", "ie": "Ireland", "gr": "Greece", "ua": "Ukraine", "rs": "Serbia", "tr": "Turkey", "il": "Israel",
    "ca": "Canada", "mx": "Mexico", "br": "Brazil", "ar": "Argentina", "cn": "China", "tw": "Taiwan", "jp": "Japan",
    "kr": "South Korea", "sg": "Singapore", "my": "Malaysia", "vn": "Vietnam", "au": "Australia", "nz": "New Zealand",
    "za": "South Africa",
}  # fmt: skip
_COUNTRY_RE = {
    c: re.compile(r"(?<![a-z])(" + "|".join(re.escape(n) for n in names) + r")(?![a-z])", re.I)
    for c, names in COUNTRIES.items()
}


# --- records ----------------------------------------------------------------------------------------


@dataclass
class Candidate:
    name: str
    website: str = ""
    category: str = ""
    found_via: str = "seed"
    aliases: list[str] = field(default_factory=list)


@dataclass
class CompanyRow:
    name: str
    website: str = ""
    country: str = ""
    other_countries: str = ""
    focus_areas: str = ""
    contact_email: str = ""
    email_source_url: str = ""
    careers_page: str = ""
    relevant_openings: int = 0
    date_checked: str = ""
    notes: str = ""
    category: str = ""
    found_via: str = ""
    accessible: bool = True
    fits: bool = True


@dataclass
class OpeningRow:
    company: str
    job_title: str
    location: str = ""
    remote: str = ""
    experience: str = ""
    job_link: str = ""
    source: str = ""
    date_checked: str = ""


# --- pure helpers (unit tested) ---------------------------------------------------------------------


class _Links(HTMLParser):
    def __init__(self, base: str) -> None:
        super().__init__()
        self.base = base
        self.links: list[tuple[str, str]] = []
        self._href: str | None = None
        self._text: list[str] = []

    def handle_starttag(self, tag, attrs):
        if tag == "a":
            self._href = dict(attrs).get("href") or ""
            self._text = []
        elif tag == "img" and self._href is not None:  # a logo link: its alt text names the company
            self._text.append(dict(attrs).get("alt") or "")

    def handle_data(self, data):
        if self._href is not None:
            self._text.append(data)

    def handle_endtag(self, tag):
        if tag == "a" and self._href is not None:
            url = urljoin(self.base, self._href.strip()).split("#")[0]
            if url.startswith(("http://", "https://")):
                self.links.append((url, " ".join(" ".join(self._text).split())))
            self._href = None


_BASE = re.compile(
    r"<base\s[^>]*href=[\"']([^\"']+)|<link\s[^>]*rel=[\"']canonical[\"'][^>]*href=[\"']([^\"']+)",
    re.I,
)


def page_base(html: str, url: str) -> str:
    """The address relative links resolve against: <base href> or the canonical address when it is on the same
    site (a redirect from '/careers' to '/careers/' changes where 'job-1' points), else the address asked for."""
    head_end = html.find("</head>")
    for m in _BASE.finditer(
        html[: head_end if head_end > 0 else 200_000]
    ):  # WordPress heads can be huge
        cand = urljoin(url, m.group(1) or m.group(2))
        if web.registrable_domain(cand) == web.registrable_domain(url):
            return cand
    return url


def links(html: str, base: str) -> list[tuple[str, str]]:
    """(absolute url, link text) for every link on the page, in page order, without duplicates."""
    p = _Links(page_base(html, base))
    try:
        p.feed(html)
    except Exception:
        log.warning("link parse error", extra={"ctx": {"page": base}})
    first: dict[str, str] = {}
    for u, t in p.links:
        first.setdefault(u, t)
    return list(first.items())


def focus_tags(text: str) -> list[str]:
    return [tag for tag, rx in _FOCUS.items() if rx.search(text or "")]


def _role_rank(address: str) -> int | None:
    local = address.split("@", 1)[0].lower()
    for rank, group in enumerate(_ROLE_RANK):
        if any(local.startswith(k) for k in group):
            return rank
    return None  # a person's address (or something odd): never collected


_AT = re.compile(r"\s*[\[({]\s*at\s*[\])}]\s*", re.I)
_DOT = re.compile(r"\s*[\[({]\s*dot\s*[\])}]\s*", re.I)


def company_emails(text: str, domains: set[str]) -> list[str]:
    """Company-level addresses published on the page, on the company's own domain(s), careers first.
    Uses the project's guard (contacts.emails_on_domain): an address must be written on the page. An address
    written as 'info [at] acme [dot] com' is published too, so it is read; nothing is ever constructed."""
    text = _DOT.sub(".", _AT.sub("@", text or ""))
    found = {e for d in domains if d for e in contacts.emails_on_domain(text, d)}
    ranked = [(r, e) for e in found if (r := _role_rank(e)) is not None]
    return [e for _, e in sorted(ranked)]


def same_company(url: str, domains: set[str]) -> bool:
    return web.registrable_domain(url) in domains


def is_ats(url: str) -> bool:
    host = (urlparse(url).hostname or "").lower()
    return any(host == h or host.endswith("." + h) for h in ATS_HOSTS)


def careers_link(page_links: list[tuple[str, str]], domains: set[str]) -> str:
    """The careers / jobs page linked from the site (own domain or an applicant system). "" if none."""
    for url, text in page_links:
        if (same_company(url, domains) or is_ats(url)) and (
            _CAREERS_LINK.search(text) or _CAREERS_LINK.search(urlparse(url).path)
        ):
            return url
    return ""


def info_links(page_links: list[tuple[str, str]], domains: set[str], limit: int = 4) -> list[str]:
    """Contact / about / imprint pages linked from the site (own domain only), best first."""
    out: list[str] = []
    for url, text in page_links:
        path = urlparse(url).path
        last = path.rstrip("/").rsplit("/", 1)[
            -1
        ]  # /company/staff/paul is about Paul, not the company
        if (
            same_company(url, domains)
            and not _NOT_INFO.search(path)
            and (_INFO_LINK.search(text) or _INFO_LINK.search(last))
            and url not in out
        ):
            out.append(url)
    out.sort(key=lambda u: 0 if re.search(r"contact|kontakt|impressum|imprint", u, re.I) else 1)
    return out[:limit]


def job_links(
    page_links: list[tuple[str, str]], careers_url: str, domains: set[str]
) -> list[tuple[str, str]]:
    """(url, title) of the openings listed on a careers page or job board. Navigation links are skipped."""
    out: list[tuple[str, str]] = []
    for url, link_text in page_links:
        if url.rstrip("/") == careers_url.rstrip("/") or not (
            same_company(url, domains) or is_ats(url)
        ):
            continue
        title = link_text
        if _GENERIC_TEXT.fullmatch(
            link_text.strip()
        ):  # "View details": the title is in the address
            title = title_from_url(url)
            if not _ROLE_WORD.search(title):
                continue
        # "View openings", "Job listings" are a way to the jobs, not a job
        if not (6 <= len(title) <= 140) or _NAV_TEXT.search(title) or _OPENINGS_LINK.search(title):
            continue
        if _ROLE_WORD.search(title) or _JOB_PATH.search(urlparse(url).path):
            out.append((url, title))
    return out


_GENERIC_TEXT = re.compile(
    r"(view|see|show)?\s*(job\s+)?(details?|description|more|info)|apply( now| here)?|read more|learn more|"
    r"know more|more info",
    re.I,
)
_DOCUMENT = re.compile(r"\.(pdf|docx?|odt)$", re.I)


def title_from_url(url: str) -> str:
    """'.../7.Embedded-Software-Linux-Senior-Engineer_-Module-Lead.docx' -> 'Embedded Software Linux Senior
    Engineer Module Lead'. Used when a job link only says 'View details'."""
    last = unquote(urlparse(url).path.rstrip("/").rsplit("/", 1)[-1])
    last = _DOCUMENT.sub("", last)
    last = re.sub(r"^[\d._\s-]+", "", last)  # leading list numbers: "7." "15."
    last = re.sub(r"-\d+$", "", last)  # WordPress duplicates: "...-1"
    return " ".join(re.sub(r"[-_\u2013]+", " ", last).split())


def is_relevant(title: str) -> bool:
    return bool(RELEVANT.search(title or ""))


def job_details(text: str) -> dict[str, str]:
    """Location, remote / hybrid / on-site and experience, as stated on the job page ('' when not stated)."""
    m = _LOCATION.search(text)
    loc = " ".join(m.group(1).split()).strip(" ,.-") if m else ""
    low = text.lower()
    remote = (
        "Remote"
        if re.search(r"\bfully remote\b|\bremote\b", low)
        else "Hybrid"
        if "hybrid" in low
        else "On-site"
        if re.search(r"on[- ]?site|in[- ]office", low)
        else ""
    )
    exp = ""
    for e in _EXPERIENCE.finditer(text):
        if "experience" in text[max(0, e.start() - 60) : e.end() + 60].lower():
            lo, plus, hi = e.group(1), e.group(2), e.group(3)
            exp = f"{lo}-{hi} years" if hi else f"{lo}+ years" if plus else f"{lo} years"
            break
    return {"location": loc[:80], "remote": remote, "experience": exp}


def is_closed(text: str) -> bool:
    return bool(_CLOSED.search(text or ""))


_ADDRESS = re.compile(
    r"headquarter|head office|registered (office|address)|corporate office|address|adresse|si[eè]ge|sede|sitz|"
    r"anschrift|located in|based in|\bhq\b|\b\d{5,6}\b",
    re.I,
)


def guess_country(info_texts: list[str], domain: str) -> tuple[str, list[str]]:
    """(headquarters country, other countries) from the contact / about / imprint pages: the country named most
    often is taken as the headquarters, then the site's country ending. Best effort; noted as such."""
    counts: Counter[str] = Counter()
    for text in info_texts:
        for country, rx in _COUNTRY_RE.items():
            for m in rx.finditer(
                text or ""
            ):  # an address names the headquarters; a mention may be a customer
                near = text[max(0, m.start() - 160) : m.end() + 20]
                counts[country] += 5 if _ADDRESS.search(near) else 1
    named = [c for c, n in counts.most_common() if n]
    tld = _TLD_COUNTRY.get(domain.rsplit(".", 1)[-1], "") if domain else ""
    hq = tld if tld and tld in named[:2] else (named[0] if named else tld)
    return hq, [c for c in named if c != hq][:5]


_SITE_NAME = re.compile(
    r"<meta\s[^>]*(?:property|name)=[\"'](?:og:site_name|application-name)[\"'][^>]*content=[\"']([^\"']{2,80})",
    re.I,
)
_TITLE = re.compile(r"<title[^>]*>(.*?)</title>", re.I | re.S)


def site_name(html: str, url: str) -> str:
    """The company's name as its own home page states it: og:site_name, else the part of the <title> that holds
    the site's name (usually after '|' or ' - '), else the domain."""
    if m := _SITE_NAME.search(html[:200_000]):
        return " ".join(unescape(m.group(1)).split())
    label = web.registrable_domain(url).split(".")[0]
    if m := _TITLE.search(html[:200_000]):
        parts = [
            p.strip()
            for p in re.split(r"\s[|\-\u2013\u2014:\u00b7]\s", unescape(m.group(1)))
            if p.strip()
        ]
        key = re.sub(r"[^a-z0-9]", "", label.lower())
        for p in (
            parts
        ):  # the part that matches the domain is the name: "Home | Acme Embedded" on acme.com
            if key and key[:5] in re.sub(r"[^a-z0-9]", "", p.lower()):
                return " ".join(p.split())[:80]
        if len(parts) > 1:
            return " ".join(parts[-1].split())[:80]
    return label.capitalize()


def norm_name(name: str) -> str:
    n = re.sub(r"\(.*?\)", " ", (name or "").lower())
    n = re.sub(r"\b(inc|ltd|llc|gmbh|ag|sa|sas|srl|bv|oy|ab|pvt|private|limited|corp|corporation|co|"
               r"group|technologies|technology|solutions|systems|software|engineering)\b", " ", n)  # fmt: skip
    return re.sub(r"[^a-z0-9]", "", n)


def dedupe(cands: list[Candidate]) -> list[Candidate]:
    """One entry per company: by registrable domain (aliases included) and by normalised name. First one wins."""
    seen_dom: set[str] = set()
    seen_name: set[str] = set()
    out = []
    for c in cands:
        doms = {web.registrable_domain(d) for d in [c.website, *c.aliases] if d}
        key = norm_name(c.name)
        if (doms & seen_dom) or (key and key in seen_name):
            continue
        seen_dom |= doms
        if key:
            seen_name.add(key)
        out.append(c)
    return out


def _url(site: str) -> str:
    site = site.strip()
    return site if site.startswith(("http://", "https://")) else f"https://{site}"


# --- fetching --------------------------------------------------------------------------------------


async def _get(url: str) -> tuple[str, str] | None:
    """(html, visible text) of a public page, or None (refused, blocked portal, bot check, error)."""
    html = await web.fetch_smart(url)
    if not html:
        return None
    text = web.html_to_text(html)
    if leadscore.block_reason(text):
        web.bot_blocked.add(urlparse(url).hostname or "")
        return None
    return html, text


async def research(  # noqa: PLR0912, PLR0915 - one company, step by step
    c: Candidate, today: str, max_openings: int = 10
) -> tuple[CompanyRow, list[OpeningRow]]:
    """Open the company's site, its contact / about pages and careers page, and collect what they publish."""
    home = _url(c.website)
    row = CompanyRow(
        name=c.name, website=home, date_checked=today, category=c.category, found_via=c.found_via
    )
    notes: list[str] = []
    got = await _get(home)
    if got is None:
        row.accessible = False
        host = urlparse(home).hostname or ""
        why = " (bot check)" if host in web.bot_blocked else ""
        row.notes = f"{CANNOT_ACCESS}{why}"
        return row, []
    html, text = got
    if not c.found_via.startswith(
        "seed"
    ):  # a found site: its search-result title is not the company's name
        row.name = site_name(html, home)
    domains = {web.registrable_domain(home), *(web.registrable_domain(a) for a in c.aliases)}
    home_links = links(html, home)
    pages: list[tuple[str, str, str]] = [(home, html, text)]  # (url, html, text)
    for url in info_links(home_links, domains):
        if (g := await _get(url)) is not None:
            pages.append((url, *g))  # noqa: PERF401 - one awaited fetch per page

    careers = careers_link(home_links, domains)
    for _, h, _ in pages[1:]:  # some sites link careers only from the about page
        if careers:
            break
        careers = careers_link(links(h, home), domains)
    row.careers_page = careers
    openings: list[OpeningRow] = []
    if careers:
        cg = await _get(careers)
        if cg is None:
            notes.append("careers page could not be opened")
        else:
            c_html, c_text = cg
            pages.append((careers, c_html, c_text))
            listed = job_links(links(c_html, careers), careers, domains)
            # No relevant job on the careers page itself: it may only point to the openings ("View openings",
            # an applicant system). Follow that link.
            if not any(is_relevant(t) for _, t in listed):
                board = next(
                    (
                        u
                        for u, t in links(c_html, careers)
                        if u.rstrip("/") != careers.rstrip("/")
                        and (is_ats(u) or (same_company(u, domains) and _OPENINGS_LINK.search(t)))
                    ),
                    "",
                )
                if board and (bg := await _get(board)) is not None:
                    listed += job_links(links(bg[0], board), board, domains)
            seen: set[str] = set()
            for url, title in listed:
                if len(openings) >= max_openings:
                    notes.append(
                        f"more than {max_openings} relevant openings: only the first are listed"
                    )
                    break
                if url in seen or not is_relevant(title):
                    continue
                seen.add(url)
                if _DOCUMENT.search(
                    urlparse(url).path
                ):  # a job description file listed on the fetched page
                    openings.append(OpeningRow(company=row.name, job_title=title, job_link=url,
                                               source="job description document on the careers page",
                                               date_checked=today))  # fmt: skip
                    continue
                jg = await _get(url)
                if jg is None:
                    continue  # job page not reachable: not recorded (every opening must come from a fetched page)
                if is_closed(jg[1]):
                    continue
                d = job_details(jg[1])
                openings.append(OpeningRow(company=row.name, job_title=title, location=d["location"],
                                           remote=d["remote"], experience=d["experience"], job_link=url,
                                           source="applicant system" if is_ats(url) else "company careers page",
                                           date_checked=today))  # fmt: skip
    else:
        notes.append("careers page not found on the site")

    all_text = " ".join(t for _, _, t in pages)
    tags = focus_tags(all_text)
    row.focus_areas = ", ".join(tags)
    if not tags:
        row.fits = False
        notes.append("target-domain wording not found on the fetched pages - check fit")

    for url, _, t in pages:
        emails = company_emails(t, domains)
        if emails:
            row.contact_email, row.email_source_url = emails[0], url
            break
    if not row.contact_email:
        form = next((u for u, h, _ in pages if web.has_contact_form(h)), "")
        row.contact_email = f"Not listed - contact form: {form}" if form else "Not listed"
        notes.append("no company email published on the site")

    info_texts = [t for u, _, t in pages[1:] if u != careers] or [text]
    row.country, others = guess_country(info_texts, web.registrable_domain(home))
    row.other_countries = ", ".join(others)
    row.relevant_openings = len(openings)
    row.notes = "; ".join(notes)
    return row, openings


# --- candidates --------------------------------------------------------------------------------------


def load_config(path: Path = CONFIG) -> dict:
    return yaml.safe_load(path.read_text(encoding="utf-8")) or {}


async def _resolve_website(name: str, skip: set[str]) -> str:
    """A seed without a website: the first company site a web search returns for its name."""
    for url, _ in await sources.web_search(f"{name} embedded engineering company", 8):
        d = web.registrable_domain(url)
        if d and web.is_company_site(d) and d not in skip and not web.blocked(url):
            return d
    return ""


async def candidates(cfg: dict, discover: bool = True) -> list[Candidate]:
    skip = set(cfg.get("skip_domains") or [])
    out: list[Candidate] = []
    for s in cfg.get("seeds") or []:
        c = Candidate(name=s["name"], website=s.get("website") or "", category=s.get("category", ""),
                      aliases=list(s.get("aliases") or []))  # fmt: skip
        if not c.website:
            c.website = await _resolve_website(c.name, skip)
            c.found_via = (
                "seed (website found by search)" if c.website else "seed (website not found)"
            )
        out.append(c)
    if not discover:
        return dedupe(out)
    found: list[Candidate] = []
    known = {web.registrable_domain(c.website) for c in out if c.website}

    def add(url: str, title: str, via: str) -> None:
        d = web.registrable_domain(url)
        if not d or d in skip or d in known or not web.is_company_site(d) or web.blocked(url):
            return
        known.add(d)
        name = re.split(r"\s[|\--:·]\s", title or "")[0].strip() or d.split(".")[0].capitalize()
        found.append(Candidate(name=name[:80], website=d, category="discovered", found_via=via))

    for q in cfg.get("discovery_queries") or []:
        for url, title in await sources.web_search(q, 10):
            add(url, title, f"search: {q}")
    for page in cfg.get("directory_pages") or []:
        if (g := await _get(page)) is not None:
            own = web.registrable_domain(page)
            for url, text in links(g[0], page):
                if web.registrable_domain(url) != own:
                    add(url, text, f"directory: {page}")
    found = found[: int(cfg.get("max_discovered", 40))]
    return dedupe(out + found)


# --- progress files ----------------------------------------------------------------------------------


def _flat(r) -> dict:
    return {k: ";".join(v) if isinstance(v, list) else v for k, v in asdict(r).items()}


def _write_csv(path: Path, rows: list, cls) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=[x.name for x in fields(cls)])
        w.writeheader()
        w.writerows(_flat(r) for r in rows)


def _append_csv(path: Path, rows: list, cls) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    new = not path.exists()
    with path.open("a", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=[x.name for x in fields(cls)])
        if new:
            w.writeheader()
        w.writerows(_flat(r) for r in rows)


def _read_csv(path: Path, cls) -> list:
    if not path.exists():
        return []
    out = []
    with path.open(newline="", encoding="utf-8") as f:
        for d in csv.DictReader(f):
            for x in fields(cls):
                v = d.get(x.name) or ""
                if x.type is bool:
                    d[x.name] = v == "True"
                elif x.type is int:
                    d[x.name] = int(v or 0)
                elif "list" in str(x.type):
                    d[x.name] = [a for a in v.split(";") if a]
                else:
                    d[x.name] = v
            out.append(cls(**{x.name: d[x.name] for x in fields(cls)}))
    return out


# --- verification --------------------------------------------------------------------------------------


async def verify(companies: list[CompanyRow], openings: list[OpeningRow], share: float = 0.2,
                 seed: int | None = None) -> list[dict]:  # fmt: skip
    """Re-open a random share of the emails and job links and check them against their pages. Also flags any
    email without a source URL. Returns one dict per check."""
    rnd = random.Random(seed)  # noqa: S311 - picks pages to re-check, nothing secret
    with_email = [c for c in companies if "@" in c.contact_email]
    checks: list[dict] = [{"company": c.name, "kind": "email", "item": c.contact_email, "ok": False,
                           "why": "email has no source URL"} for c in with_email if not c.email_source_url]  # fmt: skip
    for c in rnd.sample(with_email, max(1, round(len(with_email) * share))) if with_email else []:
        g = await _get(c.email_source_url) if c.email_source_url else None
        email = c.contact_email.lower()
        ok = g is not None and email in contacts.emails_on_domain(g[1], email.split("@", 1)[1])
        checks.append({"company": c.name, "kind": "email", "item": c.contact_email, "ok": ok,
                       "why": "" if ok else "not found again on its source page"})  # fmt: skip
    for o in rnd.sample(openings, max(1, round(len(openings) * share))) if openings else []:
        g = await _get(o.job_link)
        ok = g is not None and not is_closed(g[1])
        checks.append({"company": o.company, "kind": "opening", "item": o.job_link, "ok": ok,
                       "why": "" if ok else "job page not reachable or closed"})  # fmt: skip
    return checks


# --- workbook ------------------------------------------------------------------------------------------


def build_workbook(companies: list[CompanyRow], openings: list[OpeningRow], checks: list[dict],
                   path: Path) -> Path:  # fmt: skip
    """Companies, Openings and Summary sheets: bold frozen header, auto-filter, column widths, live links."""
    from openpyxl import Workbook  # noqa: PLC0415 - only this command needs it
    from openpyxl.styles import Font  # noqa: PLC0415

    wb = Workbook()
    link_font = Font(color="0563C1", underline="single")

    def sheet(ws, header: list[str], rows: list[list], link_cols: set[int]) -> None:
        ws.append(header)
        for cell in ws[1]:
            cell.font = Font(bold=True)
        for r in rows:
            ws.append(r)
        for row in ws.iter_rows(min_row=2):
            for i in link_cols:
                cell = row[i]
                if isinstance(cell.value, str) and cell.value.startswith(("http://", "https://")):
                    cell.hyperlink = cell.value
                    cell.font = link_font
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        for i, col in enumerate(ws.columns, 1):
            width = max(len(str(c.value or "")) for c in col)
            ws.column_dimensions[col[0].column_letter].width = min(
                max(10, width + 2), 60 if i > 1 else 40
            )

    ws = wb.active
    ws.title = "Companies"
    sheet(ws, ["Company", "Website", "Country (HQ)", "Other countries", "Focus areas", "Contact email",
               "Email source URL", "Careers page", "Relevant openings", "Date checked", "Notes"],
          [[c.name, c.website, c.country, c.other_countries, c.focus_areas, c.contact_email, c.email_source_url,
            c.careers_page or "Not found", c.relevant_openings, c.date_checked, c.notes] for c in companies],
          {1, 6, 7})  # fmt: skip

    by_company: dict[str, list[OpeningRow]] = {}
    for o in openings:
        by_company.setdefault(o.company, []).append(o)
    rows = []
    for c in companies:
        jobs = by_company.get(c.name, [])
        if jobs:
            rows += [[o.company, o.job_title, o.location, o.remote, o.experience, o.job_link, o.source,
                      o.date_checked] for o in jobs]  # fmt: skip
        elif c.accessible:
            rows.append([c.name, NO_OPENING, NO_OPENING, "", "", "", "", c.date_checked])
    sheet(wb.create_sheet("Openings"), ["Company", "Job title", "Location", "Remote?", "Experience", "Job link",
                                        "Source", "Date checked"], rows, {5})  # fmt: skip

    s = wb.create_sheet("Summary")
    with_jobs = sum(1 for c in companies if c.relevant_openings)
    failed = [ch for ch in checks if not ch["ok"]]
    s.append(["Total companies", len(companies)])
    s.append(["Companies with relevant openings", with_jobs])
    s.append(["Total relevant openings", len(openings)])
    s.append(
        ["Companies that could not be accessed", sum(1 for c in companies if not c.accessible)]
    )
    s.append(["Verification checks (failed / run)", f"{len(failed)} / {len(checks)}"])
    s.append([])
    s.append(["Country (HQ)", "Companies"])
    for country, n in Counter(c.country or "Unknown" for c in companies).most_common():
        s.append([country, n])
    s.append([])
    s.append(["Focus area", "Companies"])
    tags = Counter(t.strip() for c in companies for t in c.focus_areas.split(",") if t.strip())
    for tag, n in tags.most_common():
        s.append([tag, n])
    for row in s.iter_rows():
        if row[0].value in ("Country (HQ)", "Focus area") or row[0].row <= 5:
            row[0].font = Font(bold=True)
    s.column_dimensions["A"].width = 40
    s.column_dimensions["B"].width = 16
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path


# --- the command ---------------------------------------------------------------------------------------


async def run(*, discover: bool = True, limit: int | None = None, fresh: bool = False,
              out: Path = OUT) -> dict:  # fmt: skip
    """Build or resume the research list and write the workbook. Returns the final report."""
    cfg = load_config()
    today = date.today().isoformat()
    cand_file, comp_file, open_file = (out / "companies_candidates.csv", out / "companies_progress.csv",
                                       out / "openings_progress.csv")  # fmt: skip
    if fresh:
        for p in (cand_file, comp_file, open_file):
            p.unlink(missing_ok=True)

    cands = _read_csv(cand_file, Candidate)
    if not cands:
        cands = await candidates(cfg, discover)
        _write_csv(cand_file, cands, Candidate)
        log.info("candidates saved", extra={"ctx": {"count": len(cands), "file": str(cand_file)}})
    progress = _read_csv(comp_file, CompanyRow)
    done = {norm_name(r.name) for r in progress}
    done_sites = {web.registrable_domain(r.website) for r in progress if r.website}

    def is_done(
        c: Candidate,
    ) -> bool:  # a found company is renamed from its own site: match the site too
        return norm_name(c.name) in done or (
            bool(c.website) and web.registrable_domain(c.website) in done_sites
        )

    todo = [c for c in cands if not is_done(c) and c.website]
    for c in cands:
        if not c.website and not is_done(c):
            _append_csv(comp_file, [CompanyRow(name=c.name, date_checked=today, category=c.category,
                                               found_via=c.found_via, accessible=False,
                                               notes="website not found - check manually")], CompanyRow)  # fmt: skip
    if limit is not None:
        todo = todo[:limit]
    sem = asyncio.Semaphore(int(env("EMBEDDED_CONCURRENCY", "3") or 3))
    max_openings = int(cfg.get("max_openings_per_company", 10))

    async def one(c: Candidate) -> None:
        async with sem:
            try:
                row, jobs = await research(c, today, max_openings)
            except Exception as exc:  # one broken site never stops the list
                log.exception("research failed", extra={"ctx": {"company": c.name}})
                row, jobs = CompanyRow(name=c.name, website=_url(c.website), date_checked=today,
                                       category=c.category, found_via=c.found_via, accessible=False,
                                       notes=f"{CANNOT_ACCESS} ({type(exc).__name__})"), []  # fmt: skip
            _append_csv(
                comp_file, [row], CompanyRow
            )  # saved at once: an interrupted run resumes from here
            _append_csv(open_file, jobs, OpeningRow)
            log.info("company researched", extra={"ctx": {"company": c.name, "openings": len(jobs),
                                                          "email": row.contact_email[:60], "fits": row.fits}})  # fmt: skip

    await asyncio.gather(*(one(c) for c in todo))
    return await finish(out)


async def finish(out: Path = OUT, seed: int | None = None) -> dict:
    """Verification pass, workbook and report from the progress files (also used on its own: --report-only)."""
    rows = _read_csv(out / "companies_progress.csv", CompanyRow)
    cfg = load_config()
    skip = set(cfg.get("skip_domains") or [])
    # Seeds always go in the sheet (unreachable or unconfirmed ones with a note): they were named on purpose.
    # A found site goes in only if it was opened, is not on the skip list and its own pages use target-domain
    # wording; otherwise it is left out and listed in the report.
    companies = [
        r
        for r in rows
        if r.found_via.startswith("seed")
        or (r.accessible and r.fits and web.registrable_domain(r.website) not in skip)
    ]
    dropped = [r.name for r in rows if r not in companies]
    names = {c.name for c in companies}
    openings = [
        o for o in _read_csv(out / "openings_progress.csv", OpeningRow) if o.company in names
    ]
    checks = await verify(companies, openings, seed=seed)
    out.mkdir(parents=True, exist_ok=True)  # noqa: ASYNC240 - local folder, instant
    with (out / "verification.csv").open("w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=["company", "kind", "item", "ok", "why"])
        w.writeheader()
        w.writerows(checks)
    path = build_workbook(companies, openings, checks, out / "embedded_companies_and_openings.xlsx")
    return {
        "excel_file": str(path),
        "companies": len(companies),
        "companies_with_openings": sum(1 for c in companies if c.relevant_openings),
        "total_openings": len(openings),
        "could_not_access": [c.name for c in companies if not c.accessible],
        "not_verified": sorted({ch["company"] for ch in checks if not ch["ok"]}),
        "fit_not_confirmed": [c.name for c in companies if c.accessible and not c.fits],
        "left_out_not_embedded": dropped,
    }
